#!/usr/bin/env python3
"""Browser window for sap_equipment_extractor.py. The extractor is not changed."""

import os
import socket
import threading
import uuid
import webbrowser

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

import sap_equipment_extractor as engine

BASE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(BASE, "work", "sap")
PREVIEW = ["Tag Number", "Equipment", "Manufacturer", "Model Number", "Model Taken From",
           "Serial Number", "Part Number", "Tag Number - Source (Document / Page)",
           "Serial Number - Source (Document / Page)", "Review"]

app = Flask(__name__, template_folder=os.path.join(BASE, "templates"), static_folder=os.path.join(BASE, "static"))
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024 * 1024
JOBS = {}
LOCK = threading.Lock()


def pick_port(start=5070):
    for port in range(start, start + 20):
        with socket.socket() as sock:
            try:
                sock.bind(("127.0.0.1", port))
            except OSError:
                continue
            return port
    return start


def safe_name(name, used):
    name = secure_filename(os.path.basename(name or "file")) or "file"
    stem, ext = os.path.splitext(name)
    candidate = name
    n = 2
    while candidate.lower() in used:
        candidate = "%s-%d%s" % (stem, n, ext)
        n += 1
    used.add(candidate.lower())
    return candidate


def read_preview(path):
    book = load_workbook(path, read_only=True, data_only=True)
    if "SAP_Equipment" not in book.sheetnames:
        book.close()
        return []
    sheet = book["SAP_Equipment"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "") for cell in next(rows, [])]
    out = []
    for record in rows:
        if not any(record):
            continue
        item = {header[i]: ("" if record[i] is None else str(record[i])) for i in range(min(len(header), len(record)))}
        out.append({name: item.get(name, "") for name in PREVIEW})
        if len(out) >= 200:
            break
    book.close()
    return out


def public_job(job):
    total = job["total"] or 0
    done = job["done"] or 0
    return {
        "id": job["id"],
        "status": job["status"],
        "done": done,
        "total": total,
        "percent": int(done * 100 / total) if total else 0,
        "log": job["log"][-8:],
        "error": job["error"],
        "rows": job["rows"],
        "row_count": job["row_count"],
        "download": "/api/jobs/%s/download" % job["id"] if job["status"] == "ready" else "",
    }


def worker(job):
    def log(message):
        text = str(message).strip()
        if text:
            job["log"].append(text)

    def progress(done, total):
        job["done"] = done
        job["total"] = total

    try:
        out = engine.run(
            job["inputs"],
            output=job["output"],
            taglist=job["taglist"],
            mode=job["mode"],
            log=log,
            progress=progress,
        )
        if not out or not os.path.isfile(out):
            job["status"] = "error"
            job["error"] = job["error"] or "No readable PDF pages were found."
            return
        job["rows"] = read_preview(out)
        job["row_count"] = len(job["rows"])
        job["status"] = "ready"
    except Exception as exc:
        job["status"] = "error"
        job["error"] = str(exc)


@app.get("/")
def index():
    return render_template("sap.html")


@app.get("/api/meta")
def meta():
    return jsonify(tesseract=engine.tesseract_ready(), columns=PREVIEW)


@app.post("/api/jobs")
def create_job():
    uploads = request.files.getlist("files")
    taglist = request.files.get("taglist")
    mode = (request.form.get("mode") or "normal").strip().lower()
    if mode not in ("fast", "normal", "thorough"):
        mode = "normal"
    pdfs = [item for item in uploads if (item.filename or "").lower().endswith(".pdf")]
    if not pdfs:
        return jsonify(error="Add at least one PDF."), 400

    job_id = uuid.uuid4().hex[:12]
    folder = os.path.join(WORK, job_id, "in")
    os.makedirs(folder, exist_ok=True)
    used = set()
    paths = []
    for item in pdfs:
        name = safe_name(item.filename, used)
        dest = os.path.join(folder, name)
        item.save(dest)
        paths.append(dest)
    tag_path = None
    if taglist and taglist.filename:
        tag_path = os.path.join(folder, safe_name(taglist.filename, used))
        taglist.save(tag_path)

    job = {
        "id": job_id,
        "status": "reading",
        "done": 0,
        "total": 0,
        "log": [],
        "error": "",
        "rows": [],
        "row_count": 0,
        "inputs": paths,
        "taglist": tag_path,
        "mode": mode,
        "output": os.path.join(WORK, job_id, "SAP_Equipment_Extract.xlsx"),
    }
    with LOCK:
        JOBS[job_id] = job
    threading.Thread(target=worker, args=(job,), daemon=True).start()
    return jsonify(public_job(job))


@app.get("/api/jobs/<job_id>")
def job_status(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify(error="This extraction is no longer open. Start it again."), 404
    return jsonify(public_job(job))


@app.get("/api/jobs/<job_id>/download")
def download(job_id):
    job = JOBS.get(job_id)
    if not job or job["status"] != "ready" or not os.path.isfile(job["output"]):
        return jsonify(error="The Excel file is not ready yet."), 404
    return send_file(job["output"], as_attachment=True, download_name="SAP_Equipment_Extract.xlsx")


@app.errorhandler(413)
def too_large(_exc):
    return jsonify(error="That upload is larger than 2 GB."), 413


def serve():
    os.makedirs(WORK, exist_ok=True)
    port = pick_port()
    url = "http://127.0.0.1:%d" % port
    print("=" * 66)
    print("   SAP EQUIPMENT EXTRACTOR")
    print("   Open this address in your browser:")
    print("   %s" % url)
    print("   Press Ctrl+C to stop.")
    print("=" * 66)
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        from waitress import serve as waitress_serve
        waitress_serve(app, host="127.0.0.1", port=port, threads=8, channel_timeout=7200)
    except ImportError:
        app.run(host="127.0.0.1", port=port, threaded=True, use_reloader=False)


if __name__ == "__main__":
    serve()
